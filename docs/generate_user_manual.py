"""
Generate FX Trading Tracker User Manual in PDF and DOCX with embedded screenshots.
Follows the same branding pattern as generate_docs.py.
"""

import os
import io
from datetime import datetime, timezone

# PDF
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_JUSTIFY
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    PageBreak, Image, KeepTogether
)

# DOCX
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

C_NAVY = colors.HexColor("#08263e")
C_RED = colors.HexColor("#ec474e")
C_BLUE = colors.HexColor("#518dca")
C_GRAY = colors.HexColor("#f1f2f2")

NAVY_RGB = RGBColor(0x08, 0x26, 0x3e)
RED_RGB = RGBColor(0xec, 0x47, 0x4e)
BLUE_RGB = RGBColor(0x51, 0x8d, 0xca)

OUT_DIR = "/app/docs"
IMG_DIR = "/app/docs/screenshots"
GENERATED = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
DOC_VERSION = "2.0"

# Max image width for good PDF/DOCX layout
PDF_IMG_W = 170 * mm
DOCX_IMG_W = Inches(6.2)


def img_path(name):
    return os.path.join(IMG_DIR, name)


def img_exists(name):
    return os.path.isfile(img_path(name))


# ═══════════════════════════════════════════════════════════════════════════════
# PDF HELPERS
# ═══════════════════════════════════════════════════════════════════════════════

def _pdf_styles():
    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle("DocTitle", parent=ss["Title"], fontSize=24, textColor=C_NAVY, spaceAfter=6, fontName="Helvetica-Bold"))
    ss.add(ParagraphStyle("DocSub", parent=ss["Normal"], fontSize=11, textColor=C_BLUE, spaceAfter=20))
    ss.add(ParagraphStyle("H1", parent=ss["Heading1"], fontSize=16, textColor=C_NAVY, spaceBefore=20, spaceAfter=8, fontName="Helvetica-Bold"))
    ss.add(ParagraphStyle("H2", parent=ss["Heading2"], fontSize=12, textColor=C_NAVY, spaceBefore=14, spaceAfter=6, fontName="Helvetica-Bold"))
    ss.add(ParagraphStyle("H3", parent=ss["Heading3"], fontSize=10, textColor=C_BLUE, spaceBefore=10, spaceAfter=4, fontName="Helvetica-Bold"))
    ss.add(ParagraphStyle("Body", parent=ss["Normal"], fontSize=9, leading=13, alignment=TA_JUSTIFY))
    ss.add(ParagraphStyle("BodySmall", parent=ss["Normal"], fontSize=8, leading=11))
    ss.add(ParagraphStyle("Caption", parent=ss["Normal"], fontSize=7.5, leading=10, textColor=C_BLUE, alignment=TA_CENTER, spaceAfter=10, fontName="Helvetica-Oblique"))
    ss.add(ParagraphStyle("Note", parent=ss["Normal"], fontSize=8, leading=11, textColor=colors.HexColor("#555555"), leftIndent=12, borderColor=C_BLUE, borderWidth=1, borderPadding=4, fontName="Helvetica-Oblique"))
    ss.add(ParagraphStyle("CodeBlock", parent=ss["Normal"], fontSize=7.5, leading=10, fontName="Courier", backColor=C_GRAY, leftIndent=10, rightIndent=10, spaceBefore=4, spaceAfter=4))
    ss.add(ParagraphStyle("CellText", parent=ss["Normal"], fontSize=7, leading=9))
    return ss


def _tbl_style():
    return TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), C_NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTSIZE", (0, 0), (-1, 0), 7.5),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 1), (-1, -1), 7),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("ALIGN", (0, 0), (-1, 0), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cccccc")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, C_GRAY]),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ])


def _page_header(canvas, doc):
    canvas.saveState()
    canvas.setFillColor(C_NAVY)
    canvas.rect(0, doc.pagesize[1] - 10*mm, doc.pagesize[0], 10*mm, fill=True, stroke=False)
    canvas.setFillColor(colors.white)
    canvas.setFont("Helvetica-Bold", 8)
    canvas.drawString(12*mm, doc.pagesize[1] - 7*mm, "FX Trading Tracker  |  User Manual")
    canvas.setFont("Helvetica", 6)
    canvas.drawRightString(doc.pagesize[0] - 12*mm, doc.pagesize[1] - 7*mm, f"v{DOC_VERSION}  |  {GENERATED}")
    canvas.setStrokeColor(C_RED)
    canvas.setLineWidth(1.2)
    canvas.line(0, doc.pagesize[1] - 10*mm, doc.pagesize[0], doc.pagesize[1] - 10*mm)
    canvas.setFillColor(C_NAVY)
    canvas.setFont("Helvetica", 6)
    canvas.drawRightString(doc.pagesize[0] - 12*mm, 7*mm, f"Page {doc.page}")
    canvas.restoreState()


def _pdf_img(elems, filename, caption=None, ss=None):
    if img_exists(filename):
        from reportlab.lib.utils import ImageReader
        fp = img_path(filename)
        ir = ImageReader(fp)
        iw, ih = ir.getSize()
        ratio = ih / iw
        draw_w = min(PDF_IMG_W, iw)
        draw_h = draw_w * ratio
        img = Image(fp, width=draw_w, height=draw_h)
        img.hAlign = 'CENTER'
        elems.append(Spacer(1, 6))
        elems.append(img)
        if caption and ss:
            elems.append(Paragraph(caption, ss["Caption"]))
        else:
            elems.append(Spacer(1, 6))


# ═══════════════════════════════════════════════════════════════════════════════
# DOCX HELPERS
# ═══════════════════════════════════════════════════════════════════════════════

def _docx_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY_RGB
    return h


def _docx_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        from docx.oxml.ns import qn
        shading = cell._element.get_or_add_tcPr()
        sh_el = shading.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "08263e"})
        shading.append(sh_el)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(7.5)
            if ri % 2 == 1:
                from docx.oxml.ns import qn
                shading = cell._element.get_or_add_tcPr()
                sh_el = shading.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "f1f2f2"})
                shading.append(sh_el)


def _docx_img(doc, filename, caption=None):
    if img_exists(filename):
        doc.add_picture(img_path(filename), width=DOCX_IMG_W)
        last_para = doc.paragraphs[-1]
        last_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs:
                run.font.size = Pt(8)
                run.font.color.rgb = BLUE_RGB
                run.font.italic = True


def _docx_note(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.italic = True
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def _docx_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(8)
    run.font.name = "Courier New"
    from docx.oxml.ns import qn
    shading = p._element.get_or_add_pPr()
    sh_el = shading.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "f1f2f2"})
    shading.append(sh_el)


# ═══════════════════════════════════════════════════════════════════════════════
# PDF CONTENT
# ═══════════════════════════════════════════════════════════════════════════════

def user_manual_pdf():
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=16*mm, bottomMargin=14*mm, leftMargin=14*mm, rightMargin=14*mm)
    ss = _pdf_styles()
    e = []

    def h1(t): e.append(Paragraph(t, ss["H1"]))
    def h2(t): e.append(Paragraph(t, ss["H2"]))
    def h3(t): e.append(Paragraph(t, ss["H3"]))
    def body(t): e.append(Paragraph(t, ss["Body"]))
    def sp(n=6): e.append(Spacer(1, n))
    def note(t): e.append(Paragraph(t, ss["Note"]))
    def code(t): e.append(Paragraph(t, ss["CodeBlock"]))
    def img(fn, cap=None): _pdf_img(e, fn, cap, ss)
    def tbl(headers, rows, widths=None):
        data = [headers] + rows
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(_tbl_style())
        e.append(t)
        sp()

    # ── TITLE PAGE ──
    e.append(Spacer(1, 60))
    e.append(Paragraph("User Manual", ss["DocTitle"]))
    e.append(Paragraph(f"FX Trading Tracker Platform  |  Version {DOC_VERSION}", ss["DocSub"]))
    e.append(Spacer(1, 10))
    e.append(Paragraph(f"Generated: {GENERATED}", ss["BodySmall"]))
    e.append(Paragraph("Classification: CONFIDENTIAL", ss["BodySmall"]))
    e.append(PageBreak())

    # ── TABLE OF CONTENTS ──
    h1("Table of Contents")
    toc = [
        "1. Introduction", "2. Getting Started", "3. Dashboard",
        "4. My Deals (Trader)", "5. New Deal Ticket", "6. Deal Queue (Treasury)",
        "7. Reference Data Management (Admin)", "8. User Management (Admin)",
        "9. Transaction History (Admin)", "10. Audit Trail (Admin)",
        "11. Reports", "12. System Administration",
        "13. Appendix: Transfer Types & Crypto Networks",
    ]
    for t in toc:
        body(f"    {t}")
    e.append(PageBreak())

    # ── 1. INTRODUCTION ──
    h1("1. Introduction")
    body("The FX Trading Tracker is a web-based platform for managing foreign exchange deal tickets across a multi-role operational workflow. It supports the full deal lifecycle from ticket creation through treasury processing, with audit logging, settlement proof management, and comprehensive reporting.")
    sp()
    h2("1.1 User Roles")
    tbl(["Role", "Access Level"],
        [["Trader", "Create deal tickets, upload client proofs, view own deals, generate reports"],
         ["Treasury", "Review and process deals (confirm/return), upload processor proofs, generate reports"],
         ["Admin", "Manage users, reference data, audit trail, report permissions, bulk imports"],
         ["System Admin", "Full admin access plus database reset, system exports, and database statistics"]],
        widths=[70, 395])
    sp()
    h2("1.2 Key Concepts")
    body("<b>Deal Ticket</b> - An FX trade record containing client details, currencies, amounts, rates, and bank/crypto account information.")
    body("<b>Settlement Proof</b> - Documentary evidence (images or PDFs) of fund transfers uploaded by traders (client proofs) or treasury (processor proofs).")
    body("<b>Reference Data</b> - System-wide master data including FX clients, banks, currencies, counterparties, transaction types, and transfer types.")
    body("<b>Counterparty</b> - A PJL Group entity (e.g., CLSC, PJ, Verite) used instead of 'company' for FX Local, FX - Corporate Settlement, and FX-Intercompany transfers.")
    e.append(PageBreak())

    # ── 2. GETTING STARTED ──
    h1("2. Getting Started")
    h2("2.1 Logging In")
    body("Navigate to the application URL in your web browser. The Sign In page will display.")
    img("01_login.png", "Figure 2.1 - Sign In page")
    body("Enter your Email and Password, then click Sign In. Upon successful authentication, you will be redirected to the Dashboard.")
    sp()
    h2("2.2 Navigation")
    body("The left sidebar provides navigation to all available pages based on your role. Traders see Dashboard, My Deals, New Deal, and Reports. Treasury sees Dashboard, Deal Queue, and Reports. Admins see Dashboard, Reference Data, Users, Transactions, Audit Trail, and Reports.")
    sp()
    h2("2.3 Changing Your Password")
    body("Click Change Password in the bottom-left sidebar. Enter your current password and new password (minimum 4 characters), then click Update Password.")
    e.append(PageBreak())

    # ── 3. DASHBOARD ──
    h1("3. Dashboard")
    body("The Dashboard provides an at-a-glance overview of trading activity with summary cards, volume charts, and status distribution.")
    img("02_dashboard.png", "Figure 3.1 - Trader Dashboard")
    h2("3.1 Summary Cards")
    body("The top row displays: Total Deals, Pending, Confirmed, Returned, and Total Volume for the selected period.")
    sp()
    h2("3.2 Date Range Filters")
    body("Use the date buttons in the top-right corner: Today, Yesterday, 7D, 30D, YTD, or All Time.")
    sp()
    h2("3.3 Charts")
    body("Deals Over Time (bar chart) and Status Distribution (doughnut chart) provide visual summaries.")
    note("Note: Traders see only their own deal data. Admins and Treasury see all deals.")
    e.append(PageBreak())

    # ── 4. MY DEALS ──
    h1("4. My Deals (Trader)")
    body("The My Deals page is the primary workspace for traders to manage their deal tickets.")
    img("03_my_deals.png", "Figure 4.1 - My Deals list with Amount (currency label) and Last Action badges")
    h2("4.1 Deal List Columns")
    tbl(["Column", "Description"],
        [["Reference", "Unique deal reference (e.g., FX-20260803-0001)"],
         ["Client", "Client/company name"],
         ["Type", "Transaction type (Buy or Sell)"],
         ["Pair", "Currency pair (e.g., USD/PHP)"],
         ["Amount", "Original currency amount with label (e.g., 1,000 USD)"],
         ["Rate", "Exchange rate"],
         ["Deal Date", "Date of the deal"],
         ["Status", "Current status badge (pending, confirmed, returned, cancelled)"],
         ["Last Action", "Most recent action badge from deal history"]],
        widths=[65, 400])
    sp()
    h2("4.2 Filtering and Export")
    body("Click Filters to reveal filter controls (Status, Client, Currency, Date Range). Click Export CSV to download all visible deals as a CSV file.")
    sp()
    h2("4.3 Deal Detail Dialog")
    body("Click the eye icon on any deal row to open the full detail dialog.")
    img("05_deal_detail.png", "Figure 4.2 - Deal detail with conversion summary, Ours, and settlement proofs")
    body("The dialog shows all deal fields, currency conversion summary, bank/crypto account details, settlement proofs, and full deal history timeline.")
    sp()
    h2("4.4 Uploading Settlement Proofs")
    body("For pending or returned deals, click Upload under Client's Settlement in the detail dialog. Select image files (JPEG, PNG, WebP, GIF) or PDFs (max 10MB each).")
    note("Client settlement proofs are not applicable for FX Bank Deal transactions.")
    sp()
    h2("4.5 Cancelling a Deal")
    body("For pending deals only: open the detail dialog, click Cancel / Recall, enter a cancellation reason, and click Confirm Cancellation.")
    sp()
    h2("4.6 Editing a Returned Deal")
    body("When treasury returns a deal, a red alert shows the return remarks. Click Edit Deal to modify fields, then Save Changes or Save and Resubmit. Or click Resubmit to send as-is.")
    e.append(PageBreak())

    # ── 5. NEW DEAL TICKET ──
    h1("5. New Deal Ticket")
    body("Create a new FX trade deal ticket from the New Deal page.")
    img("04_new_deal.png", "Figure 5.1 - New Deal Ticket form")
    h2("5.1 Deal Information")
    tbl(["Field", "Required", "Description"],
        [["Client Name", "Yes", "Free-text client name"],
         ["Transaction Type", "Yes", "Buy or Sell"],
         ["Transfer Type", "Yes", "See Appendix for all types"],
         ["Deal Date", "Yes", "Date of the deal"],
         ["Value Date", "Yes", "Settlement value date"]],
        widths=[85, 50, 330])
    sp()
    h2("5.2 Amounts and Currency")
    tbl(["Field", "Required", "Description"],
        [["Currency", "Yes", "Searchable dropdown (fiat, stablecoin, crypto)"],
         ["Currency Amount", "Yes", "Principal amount in selected currency"],
         ["Exchange Rate", "Yes", "FX conversion rate"],
         ["Converted Amount", "Auto", "Calculated as Currency Amount x Rate"]],
        widths=[90, 50, 325])
    sp()
    h2("5.3 Source (From) and Destination (To)")
    body("Each section supports Bank or Crypto mode, toggled via the Bank/Crypto switch.")
    body("<b>Bank Mode:</b> Company or Counterparty (searchable) + Bank (searchable) + Account (searchable autocomplete showing account name and number).")
    body("<b>Crypto Mode:</b> Company or Counterparty + Wallet Address + Network (required: SOLANA, ETHEREUM, or TRON).")
    sp()
    h2("5.4 Ours (Receiving Account)")
    body("Defines where the counterparty credits funds. Supports Bank/Crypto modes. For FX-Intercompany deals, two Ours sections appear: Selling Counterparty Ours and Buying Counterparty Ours.")
    sp()
    h2("5.5 Submitting")
    body("Fill all required fields, click Submit Deal Ticket, review the confirmation dialog, then click Confirm and Submit. The deal is assigned a unique reference and placed in pending status.")
    e.append(PageBreak())

    # ── 6. DEAL QUEUE (TREASURY) ──
    h1("6. Deal Queue (Treasury)")
    body("The Deal Queue is the treasury team's workspace for reviewing and processing FX deals.")
    img("06_treasury_queue.png", "Figure 6.1 - Treasury Deal Queue with Last Action column")
    h2("6.1 Queue Tabs")
    tbl(["Tab", "Contents"],
        [["Pending", "Deals awaiting treasury review"],
         ["Returned", "Deals returned, awaiting trader correction"],
         ["Processed", "Confirmed and cancelled deals"]],
        widths=[80, 385])
    sp()
    h2("6.2 Customizing Columns")
    body("Click the Columns button to show/hide individual columns (Reference, Client, Type, Pair, Amount, Rate, From Bank, To Bank, Deal Date, Status, Last Action).")
    sp()
    h2("6.3 Reviewing a Deal")
    body("Click Review on a pending deal. The review dialog shows complete deal details, client proofs (view-only for treasury), processor proofs (treasury can upload/manage), and deal history. Enter Treasury Remarks (required), then click Confirm (green) or Return (red).")
    sp()
    h2("6.4 Confirmation Requirements")
    note("At least one settlement proof (client or processor) must exist before a deal can be confirmed. The Confirm button is disabled when zero proofs are uploaded.")
    e.append(PageBreak())

    # ── 7. REFERENCE DATA ──
    h1("7. Reference Data Management (Admin)")
    body("Admins manage all system reference data from the Reference Data page.")
    img("07_reference_data.png", "Figure 7.1 - Reference Data page (FX Client tab) with Import CSV/Excel button")
    h2("7.1 Data Categories")
    tbl(["Tab", "Description"],
        [["FX Client", "Client/company entities"],
         ["Banks", "Banking institutions with SWIFT codes and account management"],
         ["Counterparties", "PJL Group entities (CLSC, PJ, Verite, etc.)"],
         ["Transaction Types", "Buy / Sell"],
         ["Transfer Types", "FX Crypto, FX Local, FX Bank Deal, FX-Intercompany, FX - Corporate Settlement, PDAX WD"],
         ["Currencies", "Fiat, stablecoin, and cryptocurrency codes"]],
        widths=[85, 380])
    sp()
    h2("7.2 Adding, Editing, Deleting")
    body("Click Add New to create items. Click the pencil icon to edit or trash icon to delete (with confirmation). Fields vary by tab: Name/Code for all, SWIFT Code for Banks, Type/Symbol for Currencies.")
    sp()
    h2("7.3 Bulk Import - FX Clients")
    body("On the FX Client tab, click Import CSV/Excel. Upload a CSV or .xlsx file with columns: name (required), code (required), type (optional: Customer/Vendor). Duplicates (by code) are skipped.")
    code("name,code,type\nAcme Trading Ltd,ACME_TR,Customer\nPacific Holdings,PAC_H,Vendor")
    sp()
    h2("7.4 Bank Account Management")
    img("08_banks.png", "Figure 7.2 - Banks tab with Accounts button per bank")
    body("Each bank row has an Accounts button. In the dialog, add accounts with Account Name (required) and Account Number (required). Click Import to bulk import from CSV/Excel with columns: account_number, account_name.")
    code("account_number,account_name\n001-234-567,USD Operating Account\n001-234-568,EUR Settlement Account")
    sp()
    h2("7.5 Currencies")
    img("09_currencies.png", "Figure 7.3 - Currencies tab (fiat, stablecoin, crypto)")
    body("Pre-loaded with 31 fiat, 16 stablecoin, and 18 crypto currencies. Each has a type badge and symbol.")
    e.append(PageBreak())

    # ── 8. USER MANAGEMENT ──
    h1("8. User Management (Admin)")
    img("10_users.png", "Figure 8.1 - User Management page")
    body("View all system users with name, email, role badge, and status. Use the search bar to find users. Click Add User to create new accounts (Trader, Treasury, or Admin role). Edit or delete users via the action icons.")
    e.append(PageBreak())

    # ── 9. TRANSACTION HISTORY ──
    h1("9. Transaction History (Admin)")
    img("13_transactions.png", "Figure 9.1 - Transaction History (system-wide deal view)")
    body("Provides admins a system-wide view of all deals. Shows reference, client, trader, pair, amount, rate, deal date, and status. Click the eye icon to view full details. Use Filters and Export CSV for data extraction.")
    e.append(PageBreak())

    # ── 10. AUDIT TRAIL ──
    h1("10. Audit Trail (Admin)")
    img("12_audit_trail.png", "Figure 10.1 - Audit Trail with timestamped actions")
    body("Records every significant action: deal creation/confirmation/return/cancellation, proof uploads, user management, reference data changes (including bulk imports), and password changes.")
    sp()
    tbl(["Column", "Description"],
        [["Timestamp", "Date and time of the action"],
         ["Action", "Type (Deal Created, Proof Uploaded, companies_imported, etc.)"],
         ["Entity", "What was affected (deal, user, bank_account, company)"],
         ["Reference", "Reference number or identifier"],
         ["User", "Who performed the action (name and role)"],
         ["Details", "Human-readable description"]],
        widths=[65, 400])
    body("Click Filters to filter by action type, entity type, user name, or date range.")
    e.append(PageBreak())

    # ── 11. REPORTS ──
    h1("11. Reports")
    img("11_reports.png", "Figure 11.1 - Reports selection page")
    h2("11.1 Available Reports")
    tbl(["Report", "Description", "Key Filters"],
        [["Deal Blotter", "Complete deal log with full details", "Date, status, client, currency"],
         ["Settlement", "Deals by value date with bank details and proofs", "Date, from/to bank"],
         ["Open Positions", "Pending deals by pair showing net exposure", "None (current snapshot)"],
         ["Audit Trail", "Full action history with timestamps", "Date, user name"],
         ["User Activity", "Actions per user over a period", "Date range"],
         ["Volume Summary", "Deal counts/volumes by day/week/month", "Date range, grouping"],
         ["Client Activity", "Per-client volume, frequency, avg deal size", "Date range, client"],
         ["TMS Report (SAP)", "21-column SAP mass upload format", "Date range"]],
        widths=[75, 210, 180])
    sp()
    h2("11.2 Exporting")
    body("Each report supports CSV and PDF exports. PDF documents are branded with company colors (navy headers, red accents, blue subheadings).")
    sp()
    h2("11.3 Report Permissions")
    body("Admins can control which roles access each report via the Permissions button. Toggle access for Trader, Treasury, and Admin roles per report.")
    sp()
    h2("11.4 FX-Intercompany Handling")
    body("Settlement and TMS reports generate dual lines for FX-Intercompany deals: .1 suffix for Sell side, .2 suffix for Buy side.")
    e.append(PageBreak())

    # ── 12. SYSTEM ADMINISTRATION ──
    h1("12. System Administration")
    note("The System page is only accessible to users with the sysadmin role.")
    sp()
    h2("12.1 Database Statistics")
    body("View record counts for all collections (deals, users, audit logs, reference data, etc.).")
    sp()
    h2("12.2 Data Export")
    body("Export any collection as CSV or JSON: Deals, Audit Logs, Users, Companies, Banks, Bank Accounts, Currencies, Transaction Types, Transfer Types, Counterparties.")
    sp()
    h2("12.3 Database Reset")
    body("Perform a controlled reset: Cleared = Deals, audit logs, counters, report permissions. Retained = Users, reference data, bank accounts. Type 'RESET DATABASE' to confirm.")
    note("WARNING: This action is irreversible. All deal and audit data will be permanently deleted.")
    e.append(PageBreak())

    # ── 13. APPENDIX ──
    h1("13. Appendix: Transfer Types & Crypto Networks")
    h2("Transfer Types Reference")
    tbl(["Transfer Type", "Code", "Counterparty", "Destination", "Notes"],
        [["FX Crypto Conversion", "FX_CRYPTO", "No", "Yes", "Standard crypto conversion"],
         ["FX Local", "FX_LOCAL", "Yes", "Yes", "Shows Counterparty instead of Company"],
         ["PDAX Withdrawal", "PDAX_WD", "No", "Yes", "PDAX-specific withdrawal"],
         ["FX Bank Deal", "FX_BANK", "No", "Hidden", "Destination not applicable"],
         ["FX-Intercompany", "FX_INTERCO", "Yes", "Yes", "Dual Ours sections (Selling + Buying)"],
         ["FX - Corporate Settlement", "FX_CORP_SETTLE", "Yes", "Yes", "Behaves like FX Local"]],
        widths=[95, 70, 60, 55, 185])
    sp()
    h2("Crypto Networks")
    body("When any account section is set to Crypto mode, a Network dropdown is required:")
    tbl(["Network", "Blockchain"],
        [["SOLANA", "Solana blockchain"],
         ["ETHEREUM", "Ethereum blockchain (including ERC-20 tokens)"],
         ["TRON", "TRON blockchain (including TRC-20 tokens)"]],
        widths=[80, 385])
    sp()
    h2("Status Workflow")
    body("Created -> Pending -> Confirmed (by Treasury) OR Returned (by Treasury) OR Cancelled (by Trader). Returned deals can be edited and resubmitted back to Pending.")

    # ── BUILD ──
    doc.build(e, onFirstPage=_page_header, onLaterPages=_page_header)
    buf.seek(0)
    path = os.path.join(OUT_DIR, "FX_Trading_Tracker_User_Manual.pdf")
    with open(path, "wb") as f:
        f.write(buf.read())
    print(f"  PDF: {path}")


# ═══════════════════════════════════════════════════════════════════════════════
# DOCX CONTENT
# ═══════════════════════════════════════════════════════════════════════════════

def user_manual_docx():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    def h1(t): _docx_heading(doc, t, 1)
    def h2(t): _docx_heading(doc, t, 2)
    def h3(t): _docx_heading(doc, t, 3)
    def body(t): doc.add_paragraph(t)
    def tbl(h, r): _docx_table(doc, h, r)
    def img(fn, cap=None): _docx_img(doc, fn, cap)
    def note(t): _docx_note(doc, t)
    def code(t): _docx_code(doc, t)

    # Title
    t = doc.add_heading("User Manual", 0)
    for run in t.runs:
        run.font.color.rgb = NAVY_RGB
    sub = doc.add_paragraph(f"FX Trading Tracker Platform  |  Version {DOC_VERSION}")
    for run in sub.runs:
        run.font.color.rgb = BLUE_RGB
        run.font.size = Pt(11)
    meta = doc.add_paragraph(f"Generated: {GENERATED}\nClassification: CONFIDENTIAL")
    for run in meta.runs:
        run.font.size = Pt(8)
    doc.add_page_break()

    # 1. Introduction
    h1("1. Introduction")
    body("The FX Trading Tracker is a web-based platform for managing foreign exchange deal tickets across a multi-role operational workflow. It supports the full deal lifecycle from ticket creation through treasury processing, with audit logging, settlement proof management, and comprehensive reporting.")
    h2("1.1 User Roles")
    tbl(["Role", "Access Level"],
        [["Trader", "Create deal tickets, upload client proofs, view own deals, generate reports"],
         ["Treasury", "Review and process deals (confirm/return), upload processor proofs, generate reports"],
         ["Admin", "Manage users, reference data, audit trail, report permissions, bulk imports"],
         ["System Admin", "Full admin access plus database reset, system exports, and database statistics"]])
    h2("1.2 Key Concepts")
    body("Deal Ticket - An FX trade record containing client details, currencies, amounts, rates, and bank/crypto account information.")
    body("Settlement Proof - Documentary evidence (images or PDFs) of fund transfers uploaded by traders (client proofs) or treasury (processor proofs).")
    body("Reference Data - System-wide master data including FX clients, banks, currencies, counterparties, transaction types, and transfer types.")
    body("Counterparty - A PJL Group entity (e.g., CLSC, PJ, Verite) used instead of 'company' for FX Local, FX - Corporate Settlement, and FX-Intercompany transfers.")
    doc.add_page_break()

    # 2. Getting Started
    h1("2. Getting Started")
    h2("2.1 Logging In")
    body("Navigate to the application URL. The Sign In page will display.")
    img("01_login.png", "Figure 2.1 - Sign In page")
    body("Enter your Email and Password, then click Sign In.")
    h2("2.2 Navigation")
    body("The left sidebar provides navigation based on your role. Traders: Dashboard, My Deals, New Deal, Reports. Treasury: Dashboard, Deal Queue, Reports. Admins: Dashboard, Reference Data, Users, Transactions, Audit Trail, Reports.")
    h2("2.3 Changing Your Password")
    body("Click Change Password in the bottom-left sidebar. Enter current and new password (min 4 chars).")
    doc.add_page_break()

    # 3. Dashboard
    h1("3. Dashboard")
    body("The Dashboard provides an at-a-glance overview of trading activity.")
    img("02_dashboard.png", "Figure 3.1 - Trader Dashboard")
    body("Summary cards show Total Deals, Pending, Confirmed, Returned, and Total Volume. Date filters: Today, Yesterday, 7D, 30D, YTD, All Time. Charts: Deals Over Time (bar) and Status Distribution (doughnut).")
    note("Traders see only their own deal data. Admins and Treasury see all deals.")
    doc.add_page_break()

    # 4. My Deals
    h1("4. My Deals (Trader)")
    body("The My Deals page is the primary workspace for traders.")
    img("03_my_deals.png", "Figure 4.1 - My Deals list with Amount (currency label) and Last Action badges")
    tbl(["Column", "Description"],
        [["Reference", "Unique deal reference (e.g., FX-20260803-0001)"],
         ["Client", "Client/company name"],
         ["Type", "Transaction type (Buy or Sell)"],
         ["Pair", "Currency pair (e.g., USD/PHP)"],
         ["Amount", "Original currency amount with label (e.g., 1,000 USD)"],
         ["Rate", "Exchange rate"],
         ["Deal Date", "Date of the deal"],
         ["Status", "Status badge (pending, confirmed, returned, cancelled)"],
         ["Last Action", "Most recent action from deal history"]])
    body("Click Filters to filter by Status, Client, Currency, or Date Range. Click Export CSV to download.")
    h2("4.2 Deal Detail Dialog")
    img("05_deal_detail.png", "Figure 4.2 - Deal detail with conversion summary and settlement proofs")
    body("Shows all deal fields, conversion summary, bank/crypto details, settlement proofs, and deal history. Upload client proofs, cancel pending deals, or edit/resubmit returned deals.")
    doc.add_page_break()

    # 5. New Deal
    h1("5. New Deal Ticket")
    img("04_new_deal.png", "Figure 5.1 - New Deal Ticket form")
    tbl(["Field", "Required", "Description"],
        [["Client Name", "Yes", "Free-text client name"],
         ["Transaction Type", "Yes", "Buy or Sell"],
         ["Transfer Type", "Yes", "See Appendix for all types"],
         ["Deal Date / Value Date", "Yes", "Deal and settlement dates"],
         ["Currency", "Yes", "Searchable (fiat, stablecoin, crypto)"],
         ["Currency Amount", "Yes", "Principal amount"],
         ["Exchange Rate", "Yes", "FX rate"],
         ["Converted Amount", "Auto", "Currency Amount x Rate"]])
    body("Source (From) and Destination (To) support Bank or Crypto mode. Bank: Company/Counterparty + Bank + Account (searchable autocomplete with name + number). Crypto: Wallet Address + Network (SOLANA/ETHEREUM/TRON, required).")
    body("Ours section defines where the counterparty credits funds. For FX-Intercompany, two Ours sections appear (Selling + Buying).")
    body("Click Submit Deal Ticket, review confirmation dialog, then Confirm & Submit.")
    doc.add_page_break()

    # 6. Deal Queue
    h1("6. Deal Queue (Treasury)")
    img("06_treasury_queue.png", "Figure 6.1 - Treasury Deal Queue with Last Action column")
    tbl(["Tab", "Contents"],
        [["Pending", "Deals awaiting treasury review"],
         ["Returned", "Deals returned, awaiting trader correction"],
         ["Processed", "Confirmed and cancelled deals"]])
    body("Click Columns to customize visible columns. Click Filters to narrow by Client, Currency, From/To Bank, or Date Range.")
    body("Click Review on a pending deal. Enter Treasury Remarks (required), then Confirm or Return.")
    note("At least one settlement proof must exist before confirming. The Confirm button is disabled with zero proofs.")
    doc.add_page_break()

    # 7. Reference Data
    h1("7. Reference Data Management (Admin)")
    img("07_reference_data.png", "Figure 7.1 - Reference Data (FX Client tab) with Import button")
    tbl(["Tab", "Description"],
        [["FX Client", "Client/company entities"],
         ["Banks", "Banking institutions with SWIFT codes and account management"],
         ["Counterparties", "PJL Group entities (CLSC, PJ, Verite, etc.)"],
         ["Transaction Types", "Buy / Sell"],
         ["Transfer Types", "FX Crypto, FX Local, FX Bank Deal, FX-Intercompany, FX-Corp Settlement, PDAX WD"],
         ["Currencies", "Fiat, stablecoin, and cryptocurrency codes"]])
    body("Add/edit/delete items. Bulk import FX Clients via CSV/Excel (columns: name, code, type).")
    code("name,code,type\nAcme Trading Ltd,ACME_TR,Customer")
    h2("7.2 Bank Accounts")
    img("08_banks.png", "Figure 7.2 - Banks tab with Accounts button")
    body("Click Accounts per bank. Add accounts (Name required, Number required). Bulk import via CSV/Excel (columns: account_number, account_name).")
    code("account_number,account_name\n001-234-567,USD Operating Account")
    h2("7.3 Currencies")
    img("09_currencies.png", "Figure 7.3 - Currencies (fiat, stablecoin, crypto)")
    body("Pre-loaded with 31 fiat, 16 stablecoin, and 18 crypto currencies.")
    doc.add_page_break()

    # 8. Users
    h1("8. User Management (Admin)")
    img("10_users.png", "Figure 8.1 - User Management page")
    body("View all users with name, email, role, status. Search by name/email. Add User (Trader, Treasury, Admin). Edit/delete via icons.")
    doc.add_page_break()

    # 9. Transactions
    h1("9. Transaction History (Admin)")
    img("13_transactions.png", "Figure 9.1 - Transaction History")
    body("System-wide deal view. Shows reference, client, trader, pair, amount, rate, date, status. Click eye icon for details. Use Filters and Export CSV.")
    doc.add_page_break()

    # 10. Audit Trail
    h1("10. Audit Trail (Admin)")
    img("12_audit_trail.png", "Figure 10.1 - Audit Trail")
    body("Records every action: deal lifecycle, proofs, user management, reference changes, imports, password changes.")
    tbl(["Column", "Description"],
        [["Timestamp", "Date and time of the action"],
         ["Action", "Type (Deal Created, companies_imported, etc.)"],
         ["Entity", "What was affected (deal, user, bank_account, company)"],
         ["Reference", "Reference number or identifier"],
         ["User", "Who performed the action (name and role)"],
         ["Details", "Human-readable description"]])
    doc.add_page_break()

    # 11. Reports
    h1("11. Reports")
    img("11_reports.png", "Figure 11.1 - Reports selection page")
    tbl(["Report", "Description", "Filters"],
        [["Deal Blotter", "Complete deal log", "Date, status, client, currency"],
         ["Settlement", "Deals by value date with bank details", "Date, from/to bank"],
         ["Open Positions", "Pending deals by pair (net exposure)", "None"],
         ["Audit Trail", "Full action history", "Date, user"],
         ["User Activity", "Actions per user", "Date range"],
         ["Volume Summary", "Counts/volumes by period", "Date, grouping"],
         ["Client Activity", "Per-client stats", "Date, client"],
         ["TMS (SAP)", "21-column SAP format", "Date"]])
    body("Export as CSV or branded PDF. Admins control report access via Permissions. FX-Intercompany deals generate dual lines (.1 Sell, .2 Buy).")
    doc.add_page_break()

    # 12. System Admin
    h1("12. System Administration")
    note("System page is only accessible to sysadmin role.")
    body("Database Statistics: View record counts for all collections. Data Export: Export any collection as CSV or JSON. Database Reset: Clears deals, audit logs, counters. Retains users and reference data. Type 'RESET DATABASE' to confirm.")
    note("WARNING: Database reset is irreversible.")
    doc.add_page_break()

    # 13. Appendix
    h1("13. Appendix: Transfer Types & Crypto Networks")
    h2("Transfer Types")
    tbl(["Transfer Type", "Code", "Counterparty", "Destination", "Notes"],
        [["FX Crypto Conversion", "FX_CRYPTO", "No", "Yes", "Standard crypto conversion"],
         ["FX Local", "FX_LOCAL", "Yes", "Yes", "Shows Counterparty instead of Company"],
         ["PDAX Withdrawal", "PDAX_WD", "No", "Yes", "PDAX-specific"],
         ["FX Bank Deal", "FX_BANK", "No", "Hidden", "Destination N/A"],
         ["FX-Intercompany", "FX_INTERCO", "Yes", "Yes", "Dual Ours (Selling + Buying)"],
         ["FX - Corporate Settlement", "FX_CORP_SETTLE", "Yes", "Yes", "Behaves like FX Local"]])
    h2("Crypto Networks")
    tbl(["Network", "Blockchain"],
        [["SOLANA", "Solana blockchain"],
         ["ETHEREUM", "Ethereum (including ERC-20)"],
         ["TRON", "TRON (including TRC-20)"]])
    h2("Status Workflow")
    body("Created -> Pending -> Confirmed (Treasury) | Returned (Treasury) | Cancelled (Trader)")
    body("Returned -> Edit -> Resubmit -> Pending (cycle continues)")

    path = os.path.join(OUT_DIR, "FX_Trading_Tracker_User_Manual.docx")
    doc.save(path)
    print(f"  DOCX: {path}")


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════

if __name__ == "__main__":
    os.makedirs(OUT_DIR, exist_ok=True)
    missing = [f for f in [f"{i:02d}_{n}.png" for i, n in enumerate([
        "login", "dashboard", "my_deals", "new_deal", "deal_detail",
        "treasury_queue", "reference_data", "banks", "currencies",
        "users", "reports", "audit_trail", "transactions"], 1)] if not img_exists(f)]
    if missing:
        print(f"WARNING: Missing screenshots: {missing}")
        print(f"  Run the screenshot capture script first, or generate without images.")

    print("Generating User Manual...")
    user_manual_pdf()
    user_manual_docx()
    print(f"\nDone! Files in {OUT_DIR}/")
